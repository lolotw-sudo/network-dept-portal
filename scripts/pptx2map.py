#!/usr/bin/env python3
"""把 115 網路學系學習地圖 PPT 轉成課程資料。

用法：
  python3 scripts/pptx2map.py --slides 3,4 --xlsx out.xlsx   # 抽幾張出來校對
  python3 scripts/pptx2map.py --json out.json                # 全部轉成網頁用 JSON

判讀邏輯（會有誤差的部分都會標在「需確認」欄）：
  難度 = 方塊的 Y 座標落在「基礎/進階/專業」哪一條帶子
  班別 = 方塊框線顏色（藍=專案班、黑=計畫班、紅=臨時需求）
  子分類 = 與哪個區域標題方塊重疊面積最大
  其餘欄位讀自每張投影片上方的表頭表格，不靠座標，不會錯。
"""
import argparse, json, os, re, sys
from pptx import Presentation
from pptx.util import Emu

DEFAULT_XLSX = ("/Users/lolo/Library/CloudStorage/OneDrive-ChunghwaTelecomCo.,Ltd/"
                "學院-work/work庶務/網路學系portal/network-dept-portal/"
                "學習地圖PPT AI轉檔excel(for網頁顯示）.xlsx")

# 頁尾「資料來源」指向的檔案（課程內容的最終依據是這份校對表，不是 PPT）
SOURCE_URL = ("https://cht365-my.sharepoint.com/:x:/g/personal/lolo_cht_com_tw/"
              "IQBQr5lJPC7JTI25Dyt38MTvAedkWgzLJctyGPHNgeQsdHQ?e=YmDUwz")

DEFAULT_PPTX = ("/Users/lolo/Library/CloudStorage/OneDrive-ChunghwaTelecomCo.,Ltd/"
                "學院-work/work庶務/人發會&新版學習地圖/"
                "115網路學系_學習地圖_培訓師（network portal AI資料）.pptx")

LEVELS = ("基礎", "進階", "專業")
COURSE_LINE_MIN = 20000          # EMU：粗框＝課程方塊，細框＝區域標題
HEADER_MAX_H = 548640            # EMU(0.6吋)：比這矮的標題方塊＝欄位標題，往下整欄繼承
REGION_LINE_COLORS = {"99AAC5", "3978A5"}   # 裝飾用的灰藍框，不是課程
SKIP_PREFIX = ("事業單位窗口", "負責培訓師", "學習地圖", "跨域共課")

# 框線顏色 → 班別（依 PPT 說明頁圖例）
def class_type(hexcolor, theme):
    if hexcolor:
        r, g, b = int(hexcolor[0:2], 16), int(hexcolor[2:4], 16), int(hexcolor[4:6], 16)
        if r > 150 and g < 100 and b < 100:
            return "臨時需求"
        if b > 120 and b > r + 60:
            return "專案班"
        if max(r, g, b) < 80:
            return "計畫班"
        return "未判讀"
    if theme and "TEXT_1" in theme:
        return "計畫班"
    return "未判讀"

DIGITAL_RE = re.compile(r"數位教材|全\s*[EeＥ]\s*課程|[(（]\s*全\s*[EeＥ]\s*[)）]")

# 特色標示：綠底/紅底是圖例定義的，但這份 PPT 的數位教材多半直接寫在課名裡
def feature(hexcolor, name=""):
    tags = []
    if DIGITAL_RE.search(name):
        tags.append("數位教材")
    if hexcolor:
        r, g, b = int(hexcolor[0:2], 16), int(hexcolor[2:4], 16), int(hexcolor[4:6], 16)
        if not (r > 240 and g > 240 and b > 240):
            if g > r and g > b:
                tags.append("數位教材")
            elif r > g and r > b:
                tags.append("本年度新增")
            else:
                tags.append(f"未知底色#{hexcolor}")
    return "；".join(dict.fromkeys(tags)) or "—"

CTRL_RE = re.compile(r"[\x00-\x08\x0e-\x1f]")

def clean(text):
    """PPT 的軟換行(\x0b)當成換行，其餘控制字元去掉（Excel 不收）。"""
    return CTRL_RE.sub("", text.replace("\x0b", "\n").replace("\u200b", ""))


HOURS_RE = re.compile(r"[(（]\s*(\d+(?:\.\d+)?)\s*(天|HR|hr|小時)\s*[)）]")

def hours_of(name):
    m = HOURS_RE.findall(name)
    return f"{m[-1][0]}{m[-1][1]}" if m else "—"


def flatten(shapes, ox=0, oy=0, sx=1.0, sy=1.0):
    """攤平所有圖形並換算成投影片絕對座標（處理群組的座標偏移與縮放）。"""
    out = []
    for sh in shapes:
        if sh.__class__.__name__ == "GroupShape":
            xf = sh._element.grpSpPr.xfrm
            cw, ch = int(xf.chExt.cx) or 1, int(xf.chExt.cy) or 1
            nsx, nsy = sx * (sh.width or 1) / cw, sy * (sh.height or 1) / ch
            gl, gt = ox + (sh.left or 0) * sx, oy + (sh.top or 0) * sy
            out += flatten(sh.shapes, gl - int(xf.chOff.x) * nsx,
                           gt - int(xf.chOff.y) * nsy, nsx, nsy)
        else:
            out.append((sh, ox + (sh.left or 0) * sx, oy + (sh.top or 0) * sy,
                        (sh.width or 0) * sx, (sh.height or 0) * sy))
    return out


def colors_of(sh):
    line_hex = line_theme = fill_hex = None
    try:
        c = sh.line.color
        if c.type == 1:
            line_hex = str(c.rgb)
        elif c.type == 2:
            line_theme = str(c.theme_color)
    except Exception:
        pass
    try:
        if sh.fill.type == 1 and sh.fill.fore_color.type == 1:
            fill_hex = str(sh.fill.fore_color.rgb)
    except Exception:
        pass
    return line_hex, line_theme, fill_hex


def header_of(slide):
    """讀投影片上方的統一表頭表格。"""
    for sh in slide.shapes:
        if not sh.has_table or len(sh.table.columns) != 4:
            continue
        rows = [[clean(c.text).strip() for c in r.cells] for r in sh.table.rows]
        body = rows[1:] if rows and rows[0][0].startswith("學習地圖") else rows
        title = next((r[0] for r in body if r[0]), "")
        title = re.sub(r"[(（]\s*\d\s*/\s*\d\s*[)）]", " ", title)   # 「(1/2)」只是 PPT 排版分頁，不是不同地圖
        flat_title = " ".join(title.split())
        mt = re.match(r"^([\d.]+?)\.?\s+(.*)$", flat_title)            # 編號和名稱可能同一行也可能分兩行
        map_id = (mt.group(1) if mt else flat_title).rstrip(".")
        name = (mt.group(2) if mt else "").strip()
        seg = [x.strip() for x in name.split(".") if x.strip()]
        join = lambda i: "；".join(dict.fromkeys(
            r[i].replace("\n", " ") for r in body if len(r) > i and r[i].strip()))
        return {"mapId": map_id.rstrip("."), "mapName": name,
                "L1": seg[0] if seg else "", "L2": seg[1] if len(seg) > 1 else "",
                "L3": seg[2] if len(seg) > 2 else "",
                "audience": join(1), "contact": join(2), "trainer": join(3)}
    return None


NOTE_RE = re.compile(r"註\d")

def arch_of(slide):
    """讀 PPT 第 1 頁的架構圖（學習地圖 2.0），還原成總覽用的欄位結構。

    圖上的磚塊會橫跨欄位（例：「2.彙集」橫跨固網/海纜/衛星/行網），
    所以先用最窄的那些磚塊決定欄界，再算每塊各橫跨哪幾欄。
    """
    IN = 914400
    items, notes = [], []
    for sh, l, t, w, h in flatten(slide.shapes):
        if sh.has_table or not sh.has_text_frame:
            continue
        txt = " ".join(clean(sh.text_frame.text).split())
        if not txt or txt.startswith("Level") or l < 0.5 * IN:
            continue
        if txt.startswith("註"):
            notes += ["註" + n.strip() for n in txt.split("註") if n.strip()]
            continue
        band = ("L1" if 1.4 * IN < t < 1.7 * IN else
                "L2" if 2.7 * IN < t < 5.0 * IN else
                "L3" if 5.4 * IN < t < 6.4 * IN else None)
        if band:
            items.append({"band": band, "text": txt, "l": l, "r": l + w,
                          "t": t, "b": t + h})

    # 欄界：先用 Level 1 的磚塊定欄，再補上領域內再分欄的（例：EdgeAI 分兩欄）
    edges = sorted(i["l"] for i in items if i["band"] == "L1")
    for i in sorted(items, key=lambda i: i["l"]):
        if all(abs(i["l"] - e) > 0.25 * IN for e in edges):
            edges.append(i["l"])
    edges.sort()
    span = lambda i: [n for n, e in enumerate(edges)
                      if i["l"] - 0.25 * IN <= e < i["r"] - 0.25 * IN] or [0]

    def entry(i):
        no, _, label = i["text"].partition(".")
        label = NOTE_RE.sub("", label).replace(" ", "").strip() or i["text"]
        note = NOTE_RE.search(i["text"])
        return {"no": no.strip(), "label": label, "key": re.sub(r"[\s\-–—]", "", label),
                "cols": span(i),
                "note": next((n for n in notes if note and note.group(0) in n), "")}

    out = {"cols": len(edges), "notes": notes, "L1": [], "L2": [], "L3": []}
    for band in ("L1", "L2", "L3"):
        grp = [i for i in items if i["band"] == band]
        # 列界：把相近的 y 併成同一列，磚塊才能照原圖的位置排
        tops = sorted({i["t"] for i in grp})
        lines = []
        for y in tops:
            if not lines or y - lines[-1] > 0.3 * IN:
                lines.append(y)
        rowof = lambda y: max(n for n, ly in enumerate(lines) if y >= ly - 0.3 * IN)
        # 下緣要明確越過下一列的起點才算跨列，否則高一點的磚塊會被誤判成跨兩列
        rowend = lambda y: max([n for n, ly in enumerate(lines) if y >= ly + 0.3 * IN] or [0])
        res = []
        for i in grp:
            e = entry(i)
            e["row"] = rowof(i["t"])
            e["rowspan"] = max(1, rowend(i["b"]) - e["row"] + 1)
            res.append(e)
        out[band] = sorted(res, key=lambda e: (e["row"], e["cols"][0], -len(e["cols"])))
    return out


def parse_slide(slide, slide_no, slide_w, slide_h):
    hdr = header_of(slide)
    if not hdr:
        return []
    bands, regions, courses = [], [], []
    for sh, l, t, w, h in flatten(slide.shapes):
        if sh.has_table or not sh.has_text_frame:
            continue
        txt = " / ".join(x.strip() for x in clean(sh.text_frame.text).split("\n") if x.strip())
        if not txt or txt.isdigit():
            continue
        if not (-Emu(914400).emu < l < slide_w and 0 <= t < slide_h):   # 版面外的備註
            continue
        if txt in LEVELS:
            bands.append((txt, t, t + h))
            continue
        if txt.startswith(SKIP_PREFIX):
            continue
        line_hex, line_theme, fill_hex = colors_of(sh)
        try:
            lw = sh.line.width or 0
        except Exception:
            lw = 0
        if lw < COURSE_LINE_MIN or line_hex in REGION_LINE_COLORS:
            if len(txt) > 1:                      # PPT 上殘留的單一字元不是子分類
                regions.append((txt, l, t, w, h))
        else:
            courses.append((txt, l, t, w, h, line_hex, line_theme, fill_hex))

    bands.sort(key=lambda b: b[1])
    rows = []
    for txt, l, t, w, h, lhex, lth, fhex in sorted(courses, key=lambda c: (c[2], c[1])):
        cy, cx = t + h / 2, l + w / 2
        warn = []
        # 方塊常壓在兩條難度帶的交界上，取「重疊最多」的那條，比抓中心點穩
        ov = [(max(0, min(t + h, b[2]) - max(t, b[1])), b[0]) for b in bands]
        best_ov, lv = max(ov) if ov else (0, "未判讀")
        if best_ov <= 0:
            lv = min(bands, key=lambda b: abs((b[1] + b[2]) / 2 - cy))[0] if bands else "未判讀"
            warn.append("難度靠最近區帶推定")
        elif sorted(o[0] for o in ov)[-2:][0] > best_ov * 0.8 if len(ov) > 1 else False:
            warn.append("方塊壓在兩條難度帶交界")
        # 子分類：PPT 上有兩種標題——「框住整塊的區域方塊」與「貼在欄位頂端的細長標籤」。
        # 前者只認框得住的課程，後者才往下由整欄繼承，否則會跨到別條難度帶去。
        def xov(rl, rw):
            return max(0, min(l + w, rl + rw) - max(l, rl))
        boxed = [(rtxt, xov(rl, rw)) for rtxt, rl, rt, rw, rh in regions
                 if rt <= cy <= rt + rh and xov(rl, rw) > w * 0.3]
        heads = [(rtxt, rt) for rtxt, rl, rt, rw, rh in regions
                 if rh < HEADER_MAX_H and rt + rh <= cy and xov(rl, rw) > w * 0.3]
        if boxed:
            best = max(boxed, key=lambda x: x[1])[0]
        elif heads:
            best = max(heads, key=lambda x: x[1])[0]
        else:
            best = "未分類"
            if regions:                      # 整張都沒標子分類（例：1.1.1）就是本來不分類，不算待確認
                warn.append("PPT 上沒有對應的子分類標題")
        cols = {rtxt for rtxt, rl, rt, rw, rh in regions          # 整欄都被蓋住才算「跨欄」
                if rh < HEADER_MAX_H and rw and xov(rl, rw) >= rw * 0.8}
        if len(cols) > 1:
            warn.append(f"方塊橫跨 {len(cols)} 欄")
        ct = class_type(lhex, lth)
        if ct == "未判讀":
            warn.append(f"框線色無法對應班別({lhex or lth})")
        ft = feature(fhex, txt)
        if "未知底色" in ft:
            warn.append(ft)
        rows.append({
            "slide": slide_no, "mapId": hdr["mapId"], "mapName": hdr["mapName"],
            "L1": hdr["L1"], "L2": hdr["L2"], "L3": hdr["L3"],
            "course": txt, "hours": hours_of(txt), "level": lv, "classType": ct,
            "feature": ft, "subCat": best,
            "audience": hdr["audience"], "contact": hdr["contact"], "trainer": hdr["trainer"],
            "warn": "；".join(warn),
            "_xy": f"x={Emu(int(l)).inches:.2f} y={Emu(int(t)).inches:.2f}",
        })
    return rows


# 人工在 Excel 上改過的欄位，一律贏過從 PPT 座標判讀出來的值
OVERRIDABLE = {8: "level", 9: "classType", 10: "feature", 11: "subCat", 7: "hours"}


def apply_overrides(rows, path):
    """套用校對表上的人工修正。以（地圖編號＋課程名稱）對應，對不上的就略過。"""
    import openpyxl
    key = lambda mid, course: (str(mid).strip(), " ".join(str(course).split()))
    book = openpyxl.load_workbook(path, read_only=True).active
    table = {}
    for r in list(book.iter_rows(values_only=True))[1:]:
        if r and len(r) > 11 and r[6]:
            table[key(r[1], r[6])] = r
    hits, changed, notes = set(), 0, []
    for row in rows:
        k = key(row["mapId"], row["course"])
        src = table.get(k)
        if not src:
            continue
        hits.add(k)
        for col, field in OVERRIDABLE.items():
            val = ("" if src[col] is None else str(src[col])).strip()
            if val and val != str(row[field]).strip():
                row[field] = val
                row["warn"] = ""          # 人工指定的就不必再標待確認
                changed += 1
    stale = [k[1] for k in table if k not in hits]
    notes.append(f"校對表覆寫 {changed} 個欄位（{len(hits)}/{len(table)} 筆對得上）")
    if stale:
        notes.append(f"⚠️ 校對表有 {len(stale)} 筆在 PPT 裡找不到（課名可能改了）：{stale[0][:30]} …")
    return notes


def postprocess(rows):
    """跨投影片的收尾：同名地圖合併、整張沒分類的地圖不再嘮叨。"""
    notes = []
    # 同一張地圖被 PPT 拆成多頁（課程太多放不下）→ 以地圖名稱合併成一張
    for name in dict.fromkeys(r["mapName"] for r in rows):
        grp = [r for r in rows if r["mapName"] == name]
        ids = list(dict.fromkeys(r["mapId"] for r in grp))
        pages = list(dict.fromkeys(r["slide"] for r in grp))
        if len(pages) > 1:
            notes.append(f"「{name}」由第 {'、'.join(map(str, pages))} 頁合併為一張（{len(grp)} 筆）")
        if len(ids) > 1:
            notes.append(f"⚠️「{name}」各頁的地圖編號不一致：{ids} → 一律採用 {ids[0]}")
        for r in grp:
            r["mapId"] = ids[0]
    # 整張地圖幾乎都沒標子分類（例：1.1.1）→ 當作這張本來就不分類
    for mid in dict.fromkeys(r["mapId"] for r in rows):
        grp = [r for r in rows if r["mapId"] == mid]
        unclassified = [r for r in grp if r["subCat"] == "未分類"]
        if len(unclassified) > len(grp) / 2:
            notes.append(f"「{grp[0]['mapName']}」PPT 上未標子分類 → 整張不分類（{len(grp)} 筆）")
            for r in grp:
                r["subCat"] = "未分類"
                r["warn"] = "；".join(w for w in r["warn"].split("；")
                                      if w and "子分類" not in w)
    return notes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", default=DEFAULT_PPTX)
    ap.add_argument("--slides", help="只處理這幾張，例如 3,4")
    ap.add_argument("--xlsx", help="輸出校對用 Excel")
    ap.add_argument("--json", help="輸出網頁用 JSON")
    ap.add_argument("--overrides", default=DEFAULT_XLSX,
                    help="校對表 xlsx，上面的人工修正會蓋過 PPT 判讀結果")
    ap.add_argument("--no-overrides", action="store_true", help="忽略校對表，只用 PPT")
    ap.add_argument("--inject", action="store_true",
                    help="把資料寫回 public/ 與 dist/ 的 course-map.html")
    a = ap.parse_args()

    prs = Presentation(a.pptx)
    want = {int(x) for x in a.slides.split(",")} if a.slides else None
    rows = []
    for i, s in enumerate(prs.slides, 1):
        if want and i not in want:
            continue
        rows += parse_slide(s, i, prs.slide_width, prs.slide_height)
    notes = postprocess(rows)
    if not a.no_overrides and a.overrides and os.path.exists(a.overrides):
        notes += apply_overrides(rows, a.overrides)
    elif not a.no_overrides:
        notes.append(f"（找不到校對表，只用 PPT 判讀）：{a.overrides}")

    print(f"來源：{a.pptx}")
    print(f"抽出 {len(rows)} 筆課程，其中 {sum(1 for r in rows if r['warn'])} 筆需人工確認")
    for n in notes:
        print("  · " + n)
    for mid in dict.fromkeys(r["mapId"] for r in rows):
        sub = [r for r in rows if r["mapId"] == mid]
        print(f"  {mid} {sub[0]['mapName']}：{len(sub)} 筆"
              f"（{sum(1 for r in sub if r['warn'])} 筆待確認）")

    if a.xlsx:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "校對"
        cols = [("投影片", 8), ("地圖編號", 10), ("地圖名稱", 22), ("領域", 10), ("次領域", 10),
                ("對象", 12), ("課程名稱", 46), ("時數", 8), ("難度", 8), ("班別", 10),
                ("特色", 14), ("子分類", 18), ("⚠️需確認", 34), ("座標", 18),
                ("學習對象", 24), ("需求窗口", 16), ("培訓團隊", 34)]
        keys = ["slide", "mapId", "mapName", "L1", "L2", "L3", "course", "hours", "level",
                "classType", "feature", "subCat", "warn", "_xy", "audience", "contact", "trainer"]
        ws.append([c[0] for c in cols])
        for n, (c, wd) in enumerate(cols, 1):
            ws.cell(1, n).font = Font(bold=True)
            ws.column_dimensions[ws.cell(1, n).column_letter].width = wd
        warnfill = PatternFill("solid", fgColor="FFF2CC")
        for r in rows:
            ws.append([r[k] for k in keys])
            if r["warn"]:
                for n in range(1, len(cols) + 1):
                    ws.cell(ws.max_row, n).fill = warnfill
        ws.freeze_panes = "A2"
        wb.save(a.xlsx)
        print(f"校對表已輸出：{a.xlsx}")

    if a.inject:
        keep = ("page", "mapId", "L1", "L2", "L3", "course", "hours", "level",
                "classType", "feature", "subCat", "trainer", "contact")
        out = []
        for r in rows:
            o = dict(r, page=r["slide"])
            o["L2"] = o["L2"] or o["L1"]
            o["L3"] = o["L3"] or "通用/未標示"
            out.append({k: o[k] for k in keep})
        inject(out, a.pptx, arch_of(prs.slides[0]),
               source=a.overrides if not a.no_overrides else None)

    if a.json:
        keep = ("page", "mapId", "L1", "L2", "L3", "course", "hours", "level",
                "classType", "feature", "subCat", "trainer", "contact")
        out = []
        for r in rows:
            o = dict(r, page=r["slide"])
            o["L2"] = o["L2"] or o["L1"]              # 沒有次領域的地圖，沿用領域名當 L2
            o["L3"] = o["L3"] or "通用/未標示"          # 網頁看到這個值就不顯示 L3 標題
            out.append({k: o[k] for k in keep})
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump(out, f, ensure_ascii=False)
        print(f"JSON 已輸出：{a.json}")


def inject(rows, pptx, arch=None, source=None):
    """把課程資料寫回 public/course-map.html，並同步一份到 dist/（build 產物）。"""
    import shutil
    payload = json.dumps(rows, ensure_ascii=False)
    maps = len({r["mapId"] for r in rows})
    name = os.path.basename(source or pptx)
    src = (f'<a href="{SOURCE_URL}" target="_blank" rel="noopener">{name}</a>'
           f'（學習地圖 2.0 · {maps} 張學習地圖 · {len(rows)} 筆課程）')
    path = "public/course-map.html"
    html = open(path, encoding="utf-8").read()
    i = html.index("const RAW = [")
    j = html.index("\n", i)
    html = html[:i] + "const RAW = " + payload + ";" + html[j:]
    if arch is not None:                       # 總覽的架構直接取自 PPT 第 1 頁
        html = re.sub(r"const ARCH = .*?;\n", "const ARCH = " +
                      json.dumps(arch, ensure_ascii=False) + ";\n", html, count=1, flags=re.S)
    # 頁尾來源：內含 <a>，所以比對到第一個 </span> 為止就好
    html = re.sub(r'<span[^>]*>資料來源：.*?</span>',
                  f'<span class="src">資料來源：{src}</span>', html, count=1, flags=re.S)
    open(path, "w", encoding="utf-8").write(html)
    print(f"  已更新：{path}（{len(rows)} 筆 / {maps} 張地圖）")
    if os.path.isdir("dist"):
        shutil.copy(path, "dist/course-map.html")
        print("  已同步：dist/course-map.html")
    print("  還原用：git checkout public/course-map.html dist/course-map.html")


if __name__ == "__main__":
    sys.exit(main())
